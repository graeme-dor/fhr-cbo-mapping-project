"""Read-only verification of the published Excel distance downloads."""
import csv
from pathlib import Path
import unittest
import openpyxl

ROOT=Path(__file__).resolve().parents[1]

class ExcelDownloadTests(unittest.TestCase):
    def test_excel_preserves_all_export_values_and_numeric_distances(self):
        for filename,count,columns,numeric,table in [('distances',425,17,{3,4,13,15},'DBEDistances'),('fhr-distances',74,6,{4,5},'FHRDistances')]:
            with self.subTest(export=filename):
                folder=ROOT/'web/accessibility/data'
                with (folder/(filename+'.csv')).open(encoding='utf-8-sig',newline='') as handle:source=list(csv.reader(handle))
                workbook=openpyxl.load_workbook(folder/(filename+'.xlsx'))
                self.assertEqual(len(workbook.worksheets),1)
                sheet=workbook.active
                self.assertEqual(sheet.max_row,count+1)
                self.assertEqual(sheet.freeze_panes,'C2')
                self.assertFalse(sheet.sheet_view.showGridLines)
                self.assertEqual(len(sheet.tables),1)
                self.assertIsNotNone(sheet.tables[table].autoFilter)
                self.assertEqual(sheet.tables[table].ref,('A1:Q426' if filename=='distances' else 'A1:F75'))
                self.assertEqual([sheet.cell(1,c+1).value for c in range(columns)],source[0])
                for r,row in enumerate(source[1:],2):
                    for c,value in enumerate(row):
                        cell=sheet.cell(r,c+1)
                        self.assertNotEqual(cell.data_type,'f')
                        if c in numeric and value:
                            self.assertEqual(cell.data_type,'n')
                            self.assertAlmostEqual(cell.value,float(value),places=9)
                            self.assertEqual(cell.number_format,'0.0')
                        else:self.assertEqual(cell.value or '',value)
                workbook.close()

    def test_map_offers_excel_downloads_and_explanation_beside_selector(self):
        html=(ROOT/'web/accessibility/index.html').read_text()
        self.assertIn('data/distances.xlsx?v=',html)
        self.assertIn('data/fhr-distances.xlsx?v=',html)
        self.assertNotIn('href="data/distances.csv',html)
        self.assertNotIn('href="data/fhr-distances.csv',html)
        self.assertLess(html.index('id="grouping-help"'),html.index('<legend>Distance</legend>'))
        self.assertEqual(html.count('id="grouping-help"'),1)
        js=(ROOT/'web/accessibility/explorer.js').read_text()
        self.assertNotIn('District comparison available.',js)
        self.assertIn('District / metro matching',js)
        self.assertIn('Province-wide matching',js)

if __name__=='__main__':unittest.main()
